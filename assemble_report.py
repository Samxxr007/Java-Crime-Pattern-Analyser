import os
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

from build_report_p1 import build_report
from build_report_p2 import add_chapters_and_appendices

def assemble():
    print("Building Document Part 1 (Front matter, certificates, tables of contents)...")
    doc = build_report()
    
    print("Building Document Part 2 (Chapters 1 to 5, References, Appendices, Outcome Mapping)...")
    add_chapters_and_appendices(doc)
    
    print("Post-processing: Enforcing Times New Roman 12 black, 14pt centered main headings, and removing all lines/pBdr...")
    
    # Clean and enforce styles across all paragraphs
    for p in doc.paragraphs:
        # Remove any pBdr (paragraph borders) to ensure "no lines in page"
        pPr = p._element.pPr
        if pPr is not None:
            pBdr = pPr.find(qn('w:pBdr'))
            if pBdr is not None:
                pPr.remove(pBdr)
                
        # Check if this is a main heading
        text_strip = p.text.strip()
        is_main_heading = (
            text_strip.startswith("CHAPTER") or
            text_strip in [
                "CHENNAI CRIME PATTERN ANALYZER & RECOMMENDER",
                "BONAFIDE CERTIFICATE",
                "DECLARATION",
                "INDIVIDUAL CONTRIBUTION STATEMENT",
                "ABSTRACT",
                "TABLE OF CONTENTS",
                "LIST OF FIGURES",
                "LIST OF TABLES",
                "LIST OF ABBREVIATIONS / SYMBOLS",
                "REFERENCES",
                "APPENDICES",
                "CAPSTONE PROJECT OUTCOME MAPPING SHEET"
            ]
        )
        
        if is_main_heading:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(14)
                r.bold = True
                r.font.color.rgb = RGBColor(0, 0, 0)
                # Ensure font name is in XML
                rPr = r._r.get_or_add_rPr()
                rFonts = rPr.first_child_found_in("w:rFonts")
                if rFonts is not None:
                    rFonts.set(qn('w:ascii'), "Times New Roman")
                    rFonts.set(qn('w:hAnsi'), "Times New Roman")
        else:
            # Body text, subheadings, etc.
            for r in p.runs:
                r.font.name = "Times New Roman"
                # Keep font size 12
                r.font.size = Pt(12)
                r.font.color.rgb = RGBColor(0, 0, 0)
                rPr = r._r.get_or_add_rPr()
                rFonts = rPr.first_child_found_in("w:rFonts")
                if rFonts is not None:
                    rFonts.set(qn('w:ascii'), "Times New Roman")
                    rFonts.set(qn('w:hAnsi'), "Times New Roman")

    # Clean and enforce styles across all tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                tcPr = cell._tc.get_or_add_tcPr()
                # Ensure no weird borders inside cells except standard light single borders
                for p in cell.paragraphs:
                    pPr = p._element.pPr
                    if pPr is not None:
                        pBdr = pPr.find(qn('w:pBdr'))
                        if pBdr is not None:
                            pPr.remove(pBdr)
                    for r in p.runs:
                        r.font.name = "Times New Roman"
                        r.font.size = Pt(12)
                        r.font.color.rgb = RGBColor(0, 0, 0)
                        rPr = r._r.get_or_add_rPr()
                        rFonts = rPr.first_child_found_in("w:rFonts")
                        if rFonts is not None:
                            rFonts.set(qn('w:ascii'), "Times New Roman")
                            rFonts.set(qn('w:hAnsi'), "Times New Roman")

    out_path = "Java Capstone Report Sameer.docx"
    doc.save(out_path)
    print(f"Successfully generated and saved: {out_path}")
    print(f"Total Paragraphs: {len(doc.paragraphs)}")
    print(f"Total Tables: {len(doc.tables)}")

if __name__ == "__main__":
    assemble()
