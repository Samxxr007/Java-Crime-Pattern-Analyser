import os
import docx
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_border(cell, **kwargs):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = tcPr.first_child_found_in("w:tcBorders")
    if tcBorders is None:
        tcBorders = OxmlElement('w:tcBorders')
        tcPr.append(tcBorders)
    
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = 'w:{}'.format(edge)
            element = tcBorders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                tcBorders.append(element)
            for key, val in edge_data.items():
                element.set(qn('w:{}'.format(key)), str(val))

def set_cell_background(cell, color_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def format_run(run, font_name="Times New Roman", size_pt=12, bold=False, italic=False, color_rgb=(0, 0, 0)):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(*color_rgb)
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.first_child_found_in("w:rFonts")
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rFonts.set(qn('w:cs'), font_name)

def add_para(doc, text="", align=WD_ALIGN_PARAGRAPH.JUSTIFY, bold=False, italic=False, space_after=6, space_before=0, line_spacing=1.15):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = line_spacing
    if text:
        r = p.add_run(text)
        format_run(r, font_name="Times New Roman", size_pt=12, bold=bold, italic=italic, color_rgb=(0, 0, 0))
    return p

def add_bullet(doc, text="", bold_prefix="", space_after=4):
    p = doc.add_paragraph(style='List Bullet')
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        format_run(r_pre, font_name="Times New Roman", size_pt=12, bold=True, color_rgb=(0, 0, 0))
    if text:
        r_txt = p.add_run(text)
        format_run(r_txt, font_name="Times New Roman", size_pt=12, bold=False, color_rgb=(0, 0, 0))
    return p

def add_main_heading(doc, text, space_after=14, page_break=True):
    p = doc.add_paragraph()
    if page_break:
        p.paragraph_format.page_break_before = True
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    format_run(r, font_name="Times New Roman", size_pt=14, bold=True, italic=False, color_rgb=(0, 0, 0))
    return p

def add_sub_heading(doc, text, level=2, space_before=10, space_after=4):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    format_run(r, font_name="Times New Roman", size_pt=12, bold=True, italic=(level==3), color_rgb=(0, 0, 0))
    return p

def add_table_data(doc, headers, rows, col_widths=None, table_title=None):
    if table_title:
        add_para(doc, table_title, align=WD_ALIGN_PARAGRAPH.LEFT, bold=True, space_before=10, space_after=4)
    
    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    border_style = dict(val='single', sz='4', color='888888', space='0')
    
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], "EAECEE")
        set_cell_border(hdr_cells[i], top=border_style, bottom=border_style, left=border_style, right=border_style)
        hdr_cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.space_before = Pt(4)
        if len(p.runs) > 0:
            format_run(p.runs[0], font_name="Times New Roman", size_pt=12, bold=True, color_rgb=(0, 0, 0))
    
    for r_idx, row_data in enumerate(rows):
        row_cells = table.rows[r_idx + 1].cells
        bg_color = "F8F9F9" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, cell_value in enumerate(row_data):
            row_cells[c_idx].text = str(cell_value)
            set_cell_background(row_cells[c_idx], bg_color)
            set_cell_border(row_cells[c_idx], top=border_style, bottom=border_style, left=border_style, right=border_style)
            row_cells[c_idx].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = row_cells[c_idx].paragraphs[0]
            if c_idx == 0 or len(str(cell_value)) <= 6:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.space_before = Pt(3)
            if len(p.runs) > 0:
                format_run(p.runs[0], font_name="Times New Roman", size_pt=12, bold=False, color_rgb=(0, 0, 0))
                
    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Inches(w)
                
    add_para(doc, "", space_after=6)
    return table

def add_image_figure(doc, image_path, caption, width=Inches(6.0)):
    if os.path.exists(image_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(10)
        p_img.paragraph_format.space_after = Pt(4)
        run_img = p_img.add_run()
        run_img.add_picture(image_path, width=width)
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(2)
        p_cap.paragraph_format.space_after = Pt(10)
        run_cap = p_cap.add_run(caption)
        format_run(run_cap, font_name="Times New Roman", size_pt=12, bold=True, italic=False, color_rgb=(0, 0, 0))

print("Base setup loaded.")
